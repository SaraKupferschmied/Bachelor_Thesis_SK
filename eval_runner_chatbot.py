"""
Evaluation runner for Sara's bachelor thesis chatbot.

Reads questions from an Excel workbook, calls one or more chatbot endpoints,
and writes the returned answers/timings to a new workbook sheet.

Supports your FastAPI /ask endpoint body:
    {"question": "...", "run_mode": "rag|tool|hybrid|auto"}

Install:
    pip install openpyxl requests

Examples:
    python eval_runner_chatbot.py --input thesis_chatbot_evaluation_template.xlsx --output results.xlsx \
      --systems rag=http://localhost:8000/ask:rag hybrid=http://localhost:8000/ask:hybrid

    python eval_runner_chatbot.py --input thesis_chatbot_evaluation_template.xlsx --output test.xlsx \
      --limit 3 --systems auto=http://localhost:8000/ask:auto
"""

from __future__ import annotations

import argparse
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
import uuid

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_SHEET = "Eval_Set"
QUESTION_COLUMN_CANDIDATES = ["question", "Question", "query", "Query", "prompt", "Prompt"]

RESULT_COLUMNS = [
    "run_timestamp",
    "system",
    "run_mode",
    "http_status",
    "success",
    "client_total_time_ms",
    "server_elapsed_ms",
    "request_handler_ms",
    "answer_question_ms",
    "tool_used",
    "sources_count",
    "answer",
    "error",
]


def parse_systems(system_args: List[str]) -> Dict[str, Dict[str, Optional[str]]]:
    """
    Accepts either:
      name=url
      name=url:run_mode

    Examples:
      hybrid=http://localhost:8000/ask:hybrid
      rag=http://localhost:8000/ask:rag
      auto=http://localhost:8000/ask:auto
    """
    systems: Dict[str, Dict[str, Optional[str]]] = {}
    valid_modes = {"auto", "rag", "tool", "hybrid"}

    for item in system_args:
        if "=" not in item:
            raise ValueError(f"Invalid --systems item: {item}. Use name=url or name=url:run_mode")
        name, rest = item.split("=", 1)
        name = name.strip()
        rest = rest.strip()
        run_mode = None

        # Split a trailing :run_mode, but do not break http:// or https://.
        for mode in valid_modes:
            suffix = f":{mode}"
            if rest.endswith(suffix):
                rest = rest[: -len(suffix)]
                run_mode = mode
                break

        if not name or not rest:
            raise ValueError(f"Invalid --systems item: {item}")
        systems[name] = {"url": rest, "run_mode": run_mode}

    if not systems:
        raise ValueError("At least one system endpoint must be provided.")
    return systems


def normalize_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def find_header_row(ws) -> int:
    for row in range(1, min(ws.max_row or 1, 20) + 1):
        values = [normalize_header(ws.cell(row=row, column=col).value) for col in range(1, (ws.max_column or 1) + 1)]
        if any(v in QUESTION_COLUMN_CANDIDATES for v in values):
            return row
    raise ValueError(f"Could not find a question column. Expected one of: {QUESTION_COLUMN_CANDIDATES}")


def get_header_map(ws, header_row: int) -> Dict[str, int]:
    return {
        normalize_header(ws.cell(row=header_row, column=col).value): col
        for col in range(1, (ws.max_column or 1) + 1)
        if normalize_header(ws.cell(row=header_row, column=col).value)
    }


def find_question_col(headers: Dict[str, int]) -> int:
    for candidate in QUESTION_COLUMN_CANDIDATES:
        if candidate in headers:
            return headers[candidate]
    raise ValueError(f"Could not find a question column. Expected one of: {QUESTION_COLUMN_CANDIDATES}")


def get_measurement_ms(timing: Dict[str, Any], name: str) -> Optional[float]:
    measurements = timing.get("measurements") if isinstance(timing, dict) else None
    if not isinstance(measurements, list):
        return None
    for item in measurements:
        if isinstance(item, dict) and item.get("name") == name:
            return item.get("duration_ms")
    return None


def extract_answer(payload: Any, response_text: str) -> str:
    if isinstance(payload, dict):
        value = payload.get("answer") or payload.get("response") or payload.get("result") or payload.get("content")
        if isinstance(value, str):
            return value
        message = payload.get("message")
        if isinstance(message, dict) and isinstance(message.get("content"), str):
            return message["content"]
        if isinstance(message, str):
            return message
    return response_text


def call_system(system_name: str, url: str, question: str, run_mode: Optional[str], timeout_seconds: int) -> Dict[str, Any]:
    start = time.perf_counter()
    body: Dict[str, Any] = {
        "question": question,
        "language": "en",
        "session_id": f"eval-{system_name}-{uuid.uuid4()}",
    }
    if run_mode:
        body["run_mode"] = run_mode

    try:
        response = requests.post(url, json=body, timeout=timeout_seconds)
        client_time_ms = round((time.perf_counter() - start) * 1000, 2)
        response_text = response.text or ""

        try:
            payload: Any = response.json()
        except Exception:
            payload = None

        answer = extract_answer(payload, response_text)
        timing = payload.get("timing", {}) if isinstance(payload, dict) else {}
        used_tools = payload.get("used_tools", []) if isinstance(payload, dict) else []
        sources = payload.get("sources", []) if isinstance(payload, dict) else []

        return {
            "system": system_name,
            "run_mode": run_mode or "",
            "http_status": response.status_code,
            "success": 200 <= response.status_code < 300,
            "client_total_time_ms": client_time_ms,
            "server_elapsed_ms": timing.get("elapsed_ms") if isinstance(timing, dict) else None,
            "request_handler_ms": get_measurement_ms(timing, "request.handler") if isinstance(timing, dict) else None,
            "answer_question_ms": get_measurement_ms(timing, "ask.answer_question") if isinstance(timing, dict) else None,
            "tool_used": ", ".join(used_tools) if isinstance(used_tools, list) else str(used_tools),
            "sources_count": len(sources) if isinstance(sources, list) else None,
            "answer": answer,
            "error": None if 200 <= response.status_code < 300 else response_text[:1000],
        }
    except Exception as exc:
        client_time_ms = round((time.perf_counter() - start) * 1000, 2)
        return {
            "system": system_name,
            "run_mode": run_mode or "",
            "http_status": None,
            "success": False,
            "client_total_time_ms": client_time_ms,
            "server_elapsed_ms": None,
            "request_handler_ms": None,
            "answer_question_ms": None,
            "tool_used": "",
            "sources_count": None,
            "answer": "",
            "error": str(exc),
        }


def ensure_output_sheet(wb, base_name: str = "Auto_Run_Results"):
    if base_name in wb.sheetnames:
        del wb[base_name]
    return wb.create_sheet(base_name)


def format_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        header = str(ws.cell(row=1, column=col_idx).value or "").lower()
        if header in {"question", "answer", "expected_answer", "error"}:
            ws.column_dimensions[letter].width = 50
        elif "time" in header or "ms" in header:
            ws.column_dimensions[letter].width = 18
        else:
            ws.column_dimensions[letter].width = 16
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Path to input Excel workbook")
    parser.add_argument("--output", required=True, help="Path to output Excel workbook")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="Evaluation sheet name")
    parser.add_argument("--systems", nargs="+", required=True, help="Systems as name=url or name=url:run_mode")
    parser.add_argument("--timeout", type=int, default=180, help="Timeout per request in seconds")
    parser.add_argument("--limit", type=int, default=None, help="Only run first N questions")
    parser.add_argument("--start-row", type=int, default=None, help="Optional Excel row to start from")
    args = parser.parse_args()

    systems = parse_systems(args.systems)
    wb = load_workbook(Path(args.input))
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"Sheet '{args.sheet}' not found. Available sheets: {wb.sheetnames}")

    input_ws = wb[args.sheet]
    header_row = find_header_row(input_ws)
    headers = get_header_map(input_ws, header_row)
    question_col = find_question_col(headers)
    metadata_cols = list(range(1, input_ws.max_column + 1))

    output_ws = ensure_output_sheet(wb)
    for idx, col in enumerate(metadata_cols, start=1):
        output_ws.cell(row=1, column=idx).value = input_ws.cell(row=header_row, column=col).value
    first_result_col = len(metadata_cols) + 1

    wide_result_columns = []
    for system_name in systems.keys():
        for col in RESULT_COLUMNS:
            if col not in {"system", "run_mode"}:
                wide_result_columns.append(f"{system_name}_{col}")

    for offset, header in enumerate(wide_result_columns):
        output_ws.cell(row=1, column=first_result_col + offset).value = header

    out_row = 2
    questions_run = 0
    run_timestamp = datetime.now().isoformat(timespec="seconds")
    first_data_row = args.start_row if args.start_row is not None else header_row + 1

    for source_row in range(first_data_row, input_ws.max_row + 1):
        question = input_ws.cell(row=source_row, column=question_col).value
        if question is None or str(question).strip() == "":
            continue
        question = str(question).strip()
        if args.limit is not None and questions_run >= args.limit:
            break

        print(f"Running row {source_row}: {question[:80]}")
        for idx, col in enumerate(metadata_cols, start=1):
            output_ws.cell(row=out_row, column=idx).value = input_ws.cell(row=source_row, column=col).value

        wide_values = {}

        for system_name, spec in systems.items():
            result = call_system(system_name, spec["url"], question, spec["run_mode"], args.timeout)
            row_values = {"run_timestamp": run_timestamp, **result}

            for col in RESULT_COLUMNS:
                if col in {"system", "run_mode"}:
                    continue
                wide_values[f"{system_name}_{col}"] = row_values.get(col)

        for offset, header in enumerate(wide_result_columns):
            output_ws.cell(row=out_row, column=first_result_col + offset).value = wide_values.get(header)

        out_row += 1
        questions_run += 1

    format_sheet(output_ws)
    wb.save(Path(args.output))
    print(f"Done. Ran {questions_run} questions against {len(systems)} system(s).")
    print(f"Saved: {args.output}")


if __name__ == "__main__":
    main()
