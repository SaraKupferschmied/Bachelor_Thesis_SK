"""
Retrieval-only RAG evaluation runner for Sara's bachelor thesis chatbot.

Purpose
-------
Runs a small subset of questions through:

    POST /debug/retrieve

ONLY. It does not call /ask and therefore avoids slow LLM answer generation.

The output Excel file contains ranked retrieved sources so you can manually and
automatically compute retrieval metrics such as:

- category Recall@1 / Recall@3 / Recall@5
- category Precision@5
- category MRR
- manual Recall@K / Precision@K / MRR after annotating relevant ranks

Install
-------
    python -m pip install requests openpyxl

Example PowerShell
------------------
    python run_rag_retrieval_only_eval.py --input thesis_chatbot_evaluation_template.xlsx --output rag_retrieval_only_results.xlsx --base-url http://localhost:8000 --limit 15

After the run:
--------------
Open the output workbook and inspect retrieved_sources_ranked.
If a retrieved source is truly relevant, write its rank into manual_relevant_ranks.

Examples:
    1
    1,3
    2,4,5

Then recompute manual metrics:

    python run_rag_retrieval_only_eval.py --metrics-only --input rag_retrieval_only_results.xlsx --output rag_retrieval_only_scored.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple
import uuid

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_INPUT_SHEET = "Eval_Set"
OUTPUT_SHEET = "RAG_Retrieval_Only"
METRICS_SHEET = "Retrieval_Metrics"

QUESTION_HEADERS = ["question", "Question", "query", "Query", "prompt", "Prompt"]

RESULT_HEADERS = [
    "run_timestamp",
    "source_row",
    "id",
    "category",
    "difficulty",
    "language",
    "question",
    "expected_answer",
    "source_of_truth",
    "expected_flow",
    "gold_source",
    "required_fields",
    "expected_source_category",
    "debug_http_status",
    "success",
    "debug_time_ms",
    "server_elapsed_ms",
    "sources_count",
    "error",
    "retrieved_source_keys",
    "retrieved_source_categories",
    "retrieved_source_pages",
    "retrieved_sources_ranked",
    "retrieved_sources_json",
    "top1_source_key",
    "top1_source_category",
    "top1_source_page",
    "top1_source_snippet",
    "category_recall_at_1",
    "category_recall_at_3",
    "category_recall_at_5",
    "category_precision_at_5",
    "category_mrr",
    "manual_relevant_ranks",
    "manual_recall_at_1",
    "manual_recall_at_3",
    "manual_recall_at_5",
    "manual_precision_at_5",
    "manual_mrr",
    "notes",
]

SOURCE_CATEGORY_ALIASES = {
    "basedata": "base_data",
    "base data": "base_data",
    "base_data": "base_data",
    "baseinfo": "base_data",
    "base info": "base_data",
    "studyplan": "studyplans",
    "studyplans": "studyplans",
    "study plan": "studyplans",
    "study plans": "studyplans",
    "pdf": "studyplans",
    "curriculum": "studyplans",
    "reglementation": "reglementations",
    "reglementations": "reglementations",
    "regulation": "reglementations",
    "regulations": "reglementations",
    "rule": "reglementations",
    "rules": "reglementations",
}


def norm_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def find_header_row(ws) -> int:
    for row in range(1, min(ws.max_row or 1, 20) + 1):
        values = [norm_header(ws.cell(row=row, column=col).value) for col in range(1, ws.max_column + 1)]
        if any(v in QUESTION_HEADERS for v in values):
            return row
    raise ValueError(f"Could not find a question column. Expected one of: {QUESTION_HEADERS}")


def header_map(ws, header_row: int) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        name = norm_header(ws.cell(row=header_row, column=col).value)
        if name:
            out[name] = col
    return out


def find_question_col(headers: Dict[str, int]) -> int:
    for candidate in QUESTION_HEADERS:
        if candidate in headers:
            return headers[candidate]
    raise ValueError(f"Could not find question column. Expected one of: {QUESTION_HEADERS}")


def get_cell_by_header(ws, row: int, headers: Dict[str, int], name: str, default: Any = "") -> Any:
    col = headers.get(name)
    if col is None:
        return default
    value = ws.cell(row=row, column=col).value
    return default if value is None else value


def first_existing(row_data: Dict[str, Any], names: Sequence[str], default: Any = "") -> Any:
    for name in names:
        value = row_data.get(name)
        if value is not None and str(value).strip():
            return value
    return default


def normalize_language(value: Any) -> str:
    text = str(value or "en").strip().lower()
    if text in {"de", "ger", "german", "deutsch"}:
        return "de"
    if text in {"fr", "fre", "french", "francais", "français"}:
        return "fr"
    return "en"


def normalize_category(value: Any) -> str:
    text = str(value or "").strip().lower().replace("-", "_")
    text = re.sub(r"\s+", " ", text)
    return SOURCE_CATEGORY_ALIASES.get(text, text)


def infer_expected_source_category(row_data: Dict[str, Any]) -> str:
    joined = " ".join(str(row_data.get(k, "")) for k in [
        "question", "Question", "expected_answer", "source_of_truth", "expected_flow",
        "gold_source", "required_fields", "notes", "category"
    ]).lower()

    for needle, category in SOURCE_CATEGORY_ALIASES.items():
        if needle in joined:
            return category

    if any(word in joined for word in ["admission", "exam", "regulation", "deadline", "legal", "rule"]):
        return "reglementations"

    if any(word in joined for word in [
        "ects", "duration", "bachelor program", "master program", "doctoral program",
        "program overview", "study program", "programme overview"
    ]):
        return "base_data"

    if any(word in joined for word in [
        "mandatory", "elective", "first year", "study year", "curriculum",
        "courses include", "study plan", "module"
    ]):
        return "studyplans"

    return ""


def should_select_question(row_data: Dict[str, Any], include_all: bool = False) -> bool:
    if include_all:
        return True

    joined = " ".join(str(row_data.get(k, "")) for k in [
        "question", "Question", "source_of_truth", "expected_flow", "gold_source",
        "expected_answer", "notes", "category"
    ]).lower()

    return any(token in joined for token in [
        "rag", "mixed", "pdf", "study plan", "studyplan", "regulation",
        "reglementation", "base data", "base_data", "ects", "curriculum"
    ])


def selected_rows(
    ws,
    headers: Dict[str, int],
    start_row: int,
    limit: int,
    include_all: bool,
    ids: Optional[Sequence[str]],
) -> List[int]:
    question_col = find_question_col(headers)
    chosen: List[int] = []
    id_set = {str(x).strip() for x in ids} if ids else None

    for row in range(start_row, ws.max_row + 1):
        question = ws.cell(row=row, column=question_col).value
        if question is None or not str(question).strip():
            continue

        row_data = {name: get_cell_by_header(ws, row, headers, name, "") for name in headers}
        row_id = str(first_existing(row_data, ["id", "ID", "Id"], "")).strip()

        if id_set is not None and row_id not in id_set:
            continue
        if id_set is None and not should_select_question(row_data, include_all=include_all):
            continue

        chosen.append(row)
        if limit and len(chosen) >= limit:
            break

    if not chosen and not include_all and id_set is None:
        return selected_rows(ws, headers, start_row, limit, include_all=True, ids=None)

    return chosen


def post_json(url: str, body: Dict[str, Any], timeout: int) -> Tuple[Optional[int], Any, str, float, Optional[str]]:
    start = time.perf_counter()
    try:
        resp = requests.post(url, json=body, timeout=timeout)
        elapsed_ms = round((time.perf_counter() - start) * 1000, 2)
        text = resp.text or ""
        try:
            payload = resp.json()
        except Exception:
            payload = None
        return resp.status_code, payload, text, elapsed_ms, None
    except Exception as exc:
        elapsed_ms = round((time.perf_counter() - start) * 1000, 2)
        return None, None, "", elapsed_ms, str(exc)


def extract_sources(payload: Any) -> List[Dict[str, Any]]:
    if isinstance(payload, list):
        return [x for x in payload if isinstance(x, dict)]

    if isinstance(payload, dict):
        for key in ["sources", "documents", "results", "retrieved_documents"]:
            value = payload.get(key)
            if isinstance(value, list):
                return [x for x in value if isinstance(x, dict)]

    return []


def source_key(item: Dict[str, Any]) -> str:
    md = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
    candidates = [
        item.get("source"), item.get("source_file"), item.get("doc_key"),
        md.get("source_file"), md.get("doc_key"), md.get("chunk_id"), md.get("source_url"),
    ]
    for value in candidates:
        if value:
            return str(value)
    return "document"


def source_category(item: Dict[str, Any]) -> str:
    md = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}

    candidates = [
        md.get("rag_source"), md.get("category"), md.get("index_category"),
        md.get("source_family"), item.get("rag_source"), item.get("category"),
    ]
    for value in candidates:
        cat = normalize_category(value)
        if cat in {"base_data", "studyplans", "reglementations"}:
            return cat

    joined = " ".join(str(x or "") for x in [
        md.get("index_variant"), md.get("parser"), md.get("source_type"),
        md.get("source_file"), item.get("source"), item.get("source_type"),
    ]).lower()

    if "base" in joined:
        return "base_data"
    if "reglement" in joined or "regulation" in joined:
        return "reglementations"
    if "study" in joined or "plan" in joined or "pdf" in joined:
        return "studyplans"

    return "unknown"


def source_page(item: Dict[str, Any]) -> Any:
    md = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
    return item.get("page") or md.get("page") or md.get("page_start") or ""


def source_snippet(item: Dict[str, Any], max_chars: int = 350) -> str:
    text = str(item.get("snippet") or item.get("page_content") or item.get("content") or "")
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_chars]


def safe_json_for_excel(value: Any, max_chars: int = 32000) -> str:
    text = json.dumps(value, ensure_ascii=False, default=str)
    return text[:max_chars]


def metric_from_category(expected_category: str, categories: List[str], k: int) -> Tuple[int, float, float]:
    if not expected_category:
        return 0, 0.0, 0.0

    relevant = [1 if c == expected_category else 0 for c in categories]
    top = relevant[:k]

    recall = 1 if any(top) else 0
    precision = sum(top) / k if k else 0.0

    rr = 0.0
    for idx, is_rel in enumerate(relevant, start=1):
        if is_rel:
            rr = 1.0 / idx
            break

    return recall, precision, rr


def parse_ranks(value: Any) -> List[int]:
    if value is None:
        return []

    ranks: List[int] = []
    for part in re.split(r"[,;\s]+", str(value).strip()):
        if part.isdigit():
            ranks.append(int(part))

    return sorted(set(r for r in ranks if r > 0))


def manual_metrics_from_ranks(ranks: List[int], k: int) -> Tuple[int, float, float]:
    if not ranks:
        return 0, 0.0, 0.0

    recall = 1 if any(r <= k for r in ranks) else 0
    precision = sum(1 for r in ranks if r <= k) / k
    rr = 1.0 / min(ranks)

    return recall, precision, rr


def format_workbook(wb) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.freeze_panes = "A2"

        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            header = str(ws.cell(row=1, column=col_idx).value or "").lower()

            if "json" in header:
                ws.column_dimensions[letter].width = 45
            elif "snippet" in header or "question" in header or "ranked" in header:
                ws.column_dimensions[letter].width = 55
            elif "source" in header:
                ws.column_dimensions[letter].width = 28
            elif "time" in header or "ms" in header:
                ws.column_dimensions[letter].width = 16
            else:
                ws.column_dimensions[letter].width = 18

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_metrics_sheet(wb, result_ws_name: str = OUTPUT_SHEET) -> None:
    if METRICS_SHEET in wb.sheetnames:
        del wb[METRICS_SHEET]

    ws = wb.create_sheet(METRICS_SHEET)

    rows = [
        ["Metric", "Formula / Meaning", "Value"],
        ["Questions evaluated", "Number of rows in retrieval sheet", f"=COUNTA('{result_ws_name}'!G:G)-1"],
        ["Success rate", "Successful debug retrieval requests / all requests", f"=AVERAGE('{result_ws_name}'!O:O)"],
        ["Mean retrieval time ms", "Average /debug/retrieve client time", f"=AVERAGE('{result_ws_name}'!P:P)"],
        ["Median retrieval time ms", "Median /debug/retrieve client time", f"=MEDIAN('{result_ws_name}'!P:P)"],
        ["Mean retrieved sources count", "Average number of retrieved chunks", f"=AVERAGE('{result_ws_name}'!R:R)"],
        ["Category Recall@1", "Expected source category appears at rank 1", f"=AVERAGE('{result_ws_name}'!AC:AC)"],
        ["Category Recall@3", "Expected source category appears in top 3", f"=AVERAGE('{result_ws_name}'!AD:AD)"],
        ["Category Recall@5", "Expected source category appears in top 5", f"=AVERAGE('{result_ws_name}'!AE:AE)"],
        ["Category Precision@5", "Share of top 5 chunks with expected source category", f"=AVERAGE('{result_ws_name}'!AF:AF)"],
        ["Category MRR", "Reciprocal rank of first expected-category chunk", f"=AVERAGE('{result_ws_name}'!AG:AG)"],
        ["Manual Recall@1", "Based on manual_relevant_ranks", f"=AVERAGE('{result_ws_name}'!AI:AI)"],
        ["Manual Recall@3", "Based on manual_relevant_ranks", f"=AVERAGE('{result_ws_name}'!AJ:AJ)"],
        ["Manual Recall@5", "Based on manual_relevant_ranks", f"=AVERAGE('{result_ws_name}'!AK:AK)"],
        ["Manual Precision@5", "Based on manual_relevant_ranks", f"=AVERAGE('{result_ws_name}'!AL:AL)"],
        ["Manual MRR", "Based on manual_relevant_ranks", f"=AVERAGE('{result_ws_name}'!AM:AM)"],
    ]

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c).value = value


def run_eval(args: argparse.Namespace) -> None:
    in_path = Path(args.input)
    out_path = Path(args.output)

    base_url = args.base_url.rstrip("/")
    debug_url = f"{base_url}/debug/retrieve"

    source_wb = load_workbook(in_path)
    if args.sheet not in source_wb.sheetnames:
        raise ValueError(f"Sheet {args.sheet!r} not found. Available: {source_wb.sheetnames}")

    input_ws = source_wb[args.sheet]
    hrow = find_header_row(input_ws)
    headers = header_map(input_ws, hrow)

    rows = selected_rows(
        input_ws,
        headers,
        start_row=args.start_row or hrow + 1,
        limit=args.limit,
        include_all=args.include_all,
        ids=args.ids,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = OUTPUT_SHEET

    for col_idx, header in enumerate(RESULT_HEADERS, start=1):
        ws.cell(row=1, column=col_idx).value = header

    run_timestamp = datetime.now().isoformat(timespec="seconds")
    out_row = 2

    for source_row in rows:
        row_data = {name: get_cell_by_header(input_ws, source_row, headers, name, "") for name in headers}

        question = str(first_existing(row_data, ["question", "Question", "query", "Query", "prompt", "Prompt"], "")).strip()
        language = normalize_language(first_existing(row_data, ["language", "Language"], "en"))
        expected_category = args.expected_source_category or infer_expected_source_category(row_data)

        session_id = f"rag-retrieval-eval-{uuid.uuid4()}"

        print(f"Running row {source_row}: {question[:90]}")

        body = {
            "question": question,
            "language": language,
            "run_mode": "rag",
            "session_id": session_id,
        }

        if args.k is not None:
            body["k"] = args.k

        debug_status, debug_payload, debug_text, debug_ms, debug_err = post_json(debug_url, body, args.timeout)
        debug_sources = extract_sources(debug_payload)

        keys = [source_key(s) for s in debug_sources]
        cats = [source_category(s) for s in debug_sources]
        pages = [source_page(s) for s in debug_sources]

        ranked_text = "\n".join(
            f"{i}. [{cats[i-1]}] {keys[i-1]} p.{pages[i-1]} - {source_snippet(src, 220)}"
            for i, src in enumerate(debug_sources, start=1)
        )

        cat_r1, _, cat_mrr = metric_from_category(expected_category, cats, 1)
        cat_r3, _, _ = metric_from_category(expected_category, cats, 3)
        cat_r5, cat_p5, _ = metric_from_category(expected_category, cats, 5)

        success = int(debug_status is not None and 200 <= debug_status < 300 and len(debug_sources) > 0)

        error_parts: List[str] = []
        if debug_err:
            error_parts.append(debug_err)
        if debug_status is not None and not (200 <= debug_status < 300):
            error_parts.append(debug_text[:1000])

        values = {
            "run_timestamp": run_timestamp,
            "source_row": source_row,
            "id": first_existing(row_data, ["id", "ID", "Id"], ""),
            "category": first_existing(row_data, ["category", "Category"], ""),
            "difficulty": first_existing(row_data, ["difficulty", "Difficulty"], ""),
            "language": language,
            "question": question,
            "expected_answer": first_existing(row_data, ["expected_answer", "Expected Answer", "expected"], ""),
            "source_of_truth": first_existing(row_data, ["source_of_truth", "Source of Truth", "truth"], ""),
            "expected_flow": first_existing(row_data, ["expected_flow", "Expected Flow"], ""),
            "gold_source": first_existing(row_data, ["gold_source", "Gold Source"], ""),
            "required_fields": first_existing(row_data, ["required_fields", "Required Fields"], ""),
            "expected_source_category": expected_category,
            "debug_http_status": debug_status,
            "success": success,
            "debug_time_ms": debug_ms,
            "server_elapsed_ms": debug_payload.get("timing", {}).get("elapsed_ms") if isinstance(debug_payload, dict) and isinstance(debug_payload.get("timing"), dict) else "",
            "sources_count": len(debug_sources),
            "error": " | ".join(error_parts),
            "retrieved_source_keys": "; ".join(keys),
            "retrieved_source_categories": "; ".join(cats),
            "retrieved_source_pages": "; ".join(str(p) for p in pages),
            "retrieved_sources_ranked": ranked_text[:32000],
            "retrieved_sources_json": safe_json_for_excel(debug_sources),
            "top1_source_key": keys[0] if keys else "",
            "top1_source_category": cats[0] if cats else "",
            "top1_source_page": pages[0] if pages else "",
            "top1_source_snippet": source_snippet(debug_sources[0]) if debug_sources else "",
            "category_recall_at_1": cat_r1,
            "category_recall_at_3": cat_r3,
            "category_recall_at_5": cat_r5,
            "category_precision_at_5": cat_p5,
            "category_mrr": cat_mrr,
            "manual_relevant_ranks": "",
            "manual_recall_at_1": "",
            "manual_recall_at_3": "",
            "manual_recall_at_5": "",
            "manual_precision_at_5": "",
            "manual_mrr": "",
            "notes": "Fill manual_relevant_ranks after reviewing retrieved_sources_ranked, e.g. 1,3,5.",
        }

        for col_idx, header in enumerate(RESULT_HEADERS, start=1):
            ws.cell(row=out_row, column=col_idx).value = values.get(header, "")

        out_row += 1

        format_workbook(wb)
        write_metrics_sheet(wb)
        format_workbook(wb)
        wb.save(out_path)
        print(f"Saved progress to {out_path}")

    write_metrics_sheet(wb)
    format_workbook(wb)
    wb.save(out_path)

    print(f"Done. Wrote {len(rows)} retrieval-only tests to {out_path}")


def recompute_manual_metrics(args: argparse.Namespace) -> None:
    in_path = Path(args.input)
    out_path = Path(args.output)

    wb = load_workbook(in_path)
    if OUTPUT_SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet {OUTPUT_SHEET!r} not found in {in_path}")

    ws = wb[OUTPUT_SHEET]
    headers = {str(ws.cell(row=1, column=col).value or ""): col for col in range(1, ws.max_column + 1)}

    rank_col = headers.get("manual_relevant_ranks")
    if rank_col is None:
        raise ValueError("manual_relevant_ranks column not found")

    for row in range(2, ws.max_row + 1):
        ranks = parse_ranks(ws.cell(row=row, column=rank_col).value)

        r1, _, mrr = manual_metrics_from_ranks(ranks, 1)
        r3, _, _ = manual_metrics_from_ranks(ranks, 3)
        r5, p5, _ = manual_metrics_from_ranks(ranks, 5)

        ws.cell(row=row, column=headers["manual_recall_at_1"]).value = r1 if ranks else ""
        ws.cell(row=row, column=headers["manual_recall_at_3"]).value = r3 if ranks else ""
        ws.cell(row=row, column=headers["manual_recall_at_5"]).value = r5 if ranks else ""
        ws.cell(row=row, column=headers["manual_precision_at_5"]).value = p5 if ranks else ""
        ws.cell(row=row, column=headers["manual_mrr"]).value = mrr if ranks else ""

    write_metrics_sheet(wb)
    format_workbook(wb)
    wb.save(out_path)

    print(f"Recomputed manual metrics and wrote {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run retrieval-only RAG evaluation against /debug/retrieve.")
    parser.add_argument("--input", required=True, help="Input evaluation workbook")
    parser.add_argument("--output", required=True, help="Output workbook")
    parser.add_argument("--sheet", default=DEFAULT_INPUT_SHEET, help="Input sheet name")
    parser.add_argument("--base-url", default="http://localhost:8000", help="FastAPI base URL")
    parser.add_argument("--limit", type=int, default=15, help="Maximum number of questions to run")
    parser.add_argument("--start-row", type=int, default=None, help="Excel row to start at")
    parser.add_argument("--timeout", type=int, default=120, help="Timeout per retrieval request in seconds")
    parser.add_argument("--include-all", action="store_true", help="Select from all questions instead of RAG-focused rows")
    parser.add_argument("--ids", nargs="*", default=None, help="Optional question ids to run, e.g. --ids 4 5 12")
    parser.add_argument(
        "--expected-source-category",
        choices=["base_data", "studyplans", "reglementations"],
        default=None,
        help="Override expected source category for all selected rows",
    )
    parser.add_argument("--k", type=int, default=None, help="Optional retrieval depth, if your /debug/retrieve endpoint supports it")
    parser.add_argument("--metrics-only", action="store_true", help="Do not call chatbot; recompute manual metrics from existing output workbook")

    args = parser.parse_args()

    if args.metrics_only:
        recompute_manual_metrics(args)
    else:
        run_eval(args)


if __name__ == "__main__":
    main()
